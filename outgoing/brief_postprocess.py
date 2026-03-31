#!/usr/bin/env python3
"""
brief_postprocess.py — Post-process a WDTx brief docx after pandoc build.

Converts all Heading 1/2/3 to Normal style with paragraph-level formatting,
matching Matt's filed WDTx briefs (all Normal style, no heading styles).

Single-pass design: each paragraph is classified and formatted in one loop.

Usage:
  python3 brief_postprocess.py path/to/output.docx
"""

import sys
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, Emu
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from lxml import etree


def suppress_numpr(para):
    """Remove numPr from paragraph to suppress auto-numbering."""
    pPr = para._p.find(qn('w:pPr'))
    if pPr is None:
        return
    existing = pPr.find(qn('w:numPr'))
    if existing is not None:
        pPr.remove(existing)
    numPr = etree.SubElement(pPr, qn('w:numPr'))
    ilvl = etree.SubElement(numPr, qn('w:ilvl'))
    ilvl.set(qn('w:val'), '0')
    numId_el = etree.SubElement(numPr, qn('w:numId'))
    numId_el.set(qn('w:val'), '0')


def is_para_centered(para):
    """Check if paragraph has center alignment via XML or python-docx."""
    pPr = para._p.find(qn('w:pPr'))
    if pPr is not None:
        jc_el = pPr.find(qn('w:jc'))
        if jc_el is not None and jc_el.get(qn('w:val')) == 'center':
            return True
    return para.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER


def set_single_spaced_no_indent(para):
    """Format paragraph as single-spaced with no indent (caption/attorney/sig)."""
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.0
    pf.first_line_indent = Inches(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def set_body_formatting(para):
    """Format paragraph as double-spaced body text with 0.5" indent."""
    pf = para.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.line_spacing = 2.0
    pf.first_line_indent = Emu(457200)  # 0.5" exactly
    pf.space_before = None
    pf.space_after = None
    for run in para.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)


path = sys.argv[1]
doc = Document(path)

# ── Clear all footer text ──
cleared = 0
for section in doc.sections:
    for attr in ['footer', 'first_page_footer', 'even_page_footer']:
        try:
            f = getattr(section, attr)
            for t in f._element.iter(qn('w:t')):
                if t.text:
                    t.text = ''
                    cleared += 1
        except Exception:
            pass
print(f"  Footer: cleared {cleared} text elements")

# ── Fix table cells: single-space all tables (caption table inherits Normal=DOUBLE) ──
table_para_count = 0
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                para.paragraph_format.line_spacing = 1.0
                para.paragraph_format.first_line_indent = Inches(0)
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(0)
                table_para_count += 1
print(f"  Tables: single-spaced {table_para_count} table cell paragraphs")

# ── Single-pass paragraph formatting ──

found_first_heading = False
in_sig_block = False
past_cert_heading = False

heading_count = 0
body_count = 0
exception_count = 0

for para in doc.paragraphs:
    text = para.text.strip()
    style_name = para.style.name

    # ── Region detection ──
    if text.startswith('Dated:') or text.startswith('Respectfully submitted'):
        in_sig_block = True
    if 'CERTIFICATE OF SERVICE' in text.upper() and style_name in ('Heading 2',):
        in_sig_block = False
        past_cert_heading = True

    # ── Heading conversion (Heading 1/2/3/4 → Normal) ──
    if style_name == 'Heading 1':
        found_first_heading = True
        para.style = doc.styles['Normal']
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        para.paragraph_format.line_spacing = 2.0
        para.paragraph_format.first_line_indent = Inches(0)
        para.paragraph_format.space_before = Pt(12)
        para.paragraph_format.space_after = Pt(12)
        for run in para.runs:
            run.bold = True
        suppress_numpr(para)
        heading_count += 1
        continue

    if style_name == 'Heading 2':
        found_first_heading = True
        para.style = doc.styles['Normal']
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        para.paragraph_format.line_spacing = 2.0
        para.paragraph_format.first_line_indent = Inches(0)
        para.paragraph_format.space_before = None
        para.paragraph_format.space_after = None
        for run in para.runs:
            run.bold = True
        suppress_numpr(para)
        heading_count += 1
        continue

    if style_name == 'Heading 3':
        found_first_heading = True
        para.style = doc.styles['Normal']
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        para.paragraph_format.line_spacing = 2.0
        para.paragraph_format.first_line_indent = Inches(0)
        para.paragraph_format.space_before = None
        para.paragraph_format.space_after = None
        for run in para.runs:
            run.bold = True
        suppress_numpr(para)
        heading_count += 1
        continue

    if style_name == 'Heading 4':
        found_first_heading = True
        para.style = doc.styles['Normal']
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        para.paragraph_format.line_spacing = 2.0
        para.paragraph_format.first_line_indent = Inches(0.5)
        para.paragraph_format.space_before = None
        para.paragraph_format.space_after = None
        for run in para.runs:
            run.bold = True
        suppress_numpr(para)
        heading_count += 1
        continue

    # ── Skip non-body styles ──
    if style_name not in ('Normal', 'Body Text', 'Body Text Indent'):
        continue

    # ── Pre-heading area: caption/attorney info ──
    if not found_first_heading:
        set_single_spaced_no_indent(para)
        exception_count += 1
        continue

    # ── Signature block ──
    if in_sig_block:
        set_single_spaced_no_indent(para)
        exception_count += 1
        continue

    # ── Regular body paragraph ──
    set_body_formatting(para)
    body_count += 1

print(f"  Headings: converted {heading_count} heading paragraphs to Normal")
print(f"  Body: formatted {body_count} body paragraphs (double-spaced, 0.5\" indent)")
print(f"  Exceptions: single-spaced {exception_count} paragraphs (caption/attorney/sig)")

# ── Style-level defaults ──
# Set Normal AND Body Text styles to double-spaced TNR 12pt at the style level.
# This ensures LibreOffice renders correctly even if paragraph-level overrides
# don't take full effect (observed rendering difference between Normal and Body Text).
for style_name in ('Normal', 'Body Text'):
    try:
        style = doc.styles[style_name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        style.paragraph_format.line_spacing = 2.0
    except KeyError:
        pass

# Heading styles that may still exist in the template — override to prevent inheritance issues
for hstyle in ('Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'):
    try:
        style = doc.styles[hstyle]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(12)
    except KeyError:
        pass

print(f"  Styles: Normal/Body Text set to double-spaced TNR 12pt")

doc.save(path)
print(f"  Saved: {path}")
